import os
import ast
import json
import base64
from pathlib import Path
from email import policy
from email.parser import BytesParser
from email.message import EmailMessage
from email.utils import parseaddr

from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import Workbook, load_workbook
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build


# get the folder this file is in so it can find the other files
folder = Path(__file__).resolve().parent
load_dotenv(folder / ".env")

scopes = [
    "https://www.googleapis.com/auth/gmail.modify",
    "https://www.googleapis.com/auth/gmail.send",
]


def get_file(name):
    file1 = folder / name
    if file1.exists():
        return file1

    file2 = Path.cwd() / name
    if file2.exists():
        return file2

    return file1


# logs into gmail and returns the gmail service

def gmail_login():
    token = get_file("token.json")
    credentials = get_file("credentials.json")
    creds = None

    if token.exists():
        creds = Credentials.from_authorized_user_file(str(token), scopes)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not credentials.exists():
                raise FileNotFoundError("credentials.json was not found")

            flow = InstalledAppFlow.from_client_secrets_file(
                str(credentials), scopes
            )
            creds = flow.run_local_server(port=0)

        # save the login so i dont have to log in every time
        with open(folder / "token.json", "w", encoding="utf-8") as f:
            f.write(creds.to_json())

    return build("gmail", "v1", credentials=creds)


# gets the readable part of an email

def get_body(message):
    plain = []
    html = []

    if message.is_multipart():
        for part in message.walk():
            if part.get_content_disposition() == "attachment":
                continue

            try:
                content = part.get_content()
            except Exception:
                continue

            if part.get_content_type() == "text/plain":
                plain.append(str(content))
            elif part.get_content_type() == "text/html":
                html.append(str(content))
    else:
        try:
            content = message.get_content()
        except Exception:
            content = ""

        if message.get_content_type() == "text/plain":
            plain.append(str(content))
        elif message.get_content_type() == "text/html":
            html.append(str(content))

    if plain:
        return "\n".join(plain).strip()
    if html:
        return "\n".join(html).strip()

    return ""


# finds every unread email in the inbox

def get_unread(service):
    emails = []
    page = None

    while True:
        result = (
            service.users()
            .messages()
            .list(
                userId="me",
                q="is:unread in:inbox",
                maxResults=500,
                pageToken=page,
            )
            .execute()
        )

        emails.extend(result.get("messages", []))
        page = result.get("nextPageToken")

        if not page:
            break

    return emails


# gets the useful information from one email

def get_email(service, email_id):
    result = (
        service.users()
        .messages()
        .get(userId="me", id=email_id, format="raw")
        .execute()
    )

    decoded = base64.urlsafe_b64decode(result["raw"])
    message = BytesParser(policy=policy.default).parsebytes(decoded)

    sender = str(message.get("From", ""))
    reply_to = str(message.get("Reply-To", ""))

    return {
        "id": email_id,
        "thread_id": result.get("threadId"),
        "from": sender,
        "from_email": parseaddr(sender)[1],
        "reply_to": reply_to,
        "reply_to_email": parseaddr(reply_to)[1],
        "to": str(message.get("To", "")),
        "subject": str(message.get("Subject", "")),
        "date": str(message.get("Date", "")),
        "message_id": str(message.get("Message-ID", "")),
        "references": str(message.get("References", "")),
        "body": get_body(message),
    }


def mark_read(service, email_id):
    service.users().messages().modify(
        userId="me",
        id=email_id,
        body={"removeLabelIds": ["UNREAD"]},
    ).execute()


# sends a new email

def send_email(service, to, subject, body):
    message = EmailMessage()
    message["To"] = to
    message["Subject"] = subject
    message.set_content(body)

    encoded = base64.urlsafe_b64encode(message.as_bytes()).decode()

    return (
        service.users()
        .messages()
        .send(userId="me", body={"raw": encoded})
        .execute()
    )


# turns the ai response into a python dictionary

def get_dictionary(text):
    text = text.strip()

    # sometimes the model puts the answer inside ``` marks
    if text.startswith("```"):
        lines = text.splitlines()[1:]
        if lines and lines[-1].strip().startswith("```"):
            lines = lines[:-1]
        text = "\n".join(lines).strip()

    # only keep the actual dictionary if there is other text around it
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1:
        text = text[start:end + 1]

    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        data = ast.literal_eval(text)

    if not isinstance(data, dict):
        raise ValueError("AI did not return a dictionary")

    return data


# works out what type of email it is

def classify_email(client, email):
    prompt = (
        "Here is the data from an email: " + str(email) +
        ". Classify it as either a personal email from another person, "
        "a marketing email from a business/organisation, or a reply to an email "
        "I previously sent. I am building a tool to find organisations that may "
        "have my data and send data deletion requests. Only output one word: "
        "Personal, Marketing, or Reply."
    )

    response = client.responses.create(
        model="gpt-4.1-mini",
        input=prompt,
    )

    answer = response.output_text.strip().strip("`'\" .\n\t")

    if answer.lower() == "personal":
        return "Personal"
    if answer.lower() == "marketing":
        return "Marketing"
    if answer.lower() == "reply":
        return "Reply"

    raise ValueError("Unexpected classification: " + answer)


# makes the data request for a marketing email

def make_request(client, email):
    prompt = (
        "Here is a marketing email: " + str(email) +
        ". Find the Data Protection Office email address if one is shown. "
        "If not, use a support/contact email if there is one. If there still "
        "isn't one, use the original sender email: " + email.get("from_email", "") +
        ". Do not invent an email address. Make a subject and body for a UK data "
        "deletion request. Ask whether they store any of my personal data, where "
        "they got it from, and ask them to delete it. Only return a dictionary "
        "like this: {'email_address': 'address', 'subject': 'subject', 'body': 'body'}"
    )

    response = client.responses.create(
        model="gpt-4.1-mini",
        input=prompt,
    )

    return get_dictionary(response.output_text)


# checks whether a reply needs me to do anything else

def check_reply(client, email):
    prompt = (
        "Here is a reply to a privacy/data deletion request: " + str(email) +
        ". Work out if I need to do anything else or give more information. "
        "If they say where they got my data from, also get the name of that source. "
        "Only return a dictionary like this: "
        "{'status': 'COMPLETE' or 'FURTHER ACTION', 'datasource': 'source name or None'}"
    )

    response = client.responses.create(
        model="gpt-4.1-mini",
        input=prompt,
    )

    data = get_dictionary(response.output_text)
    status = str(data.get("status", "")).strip().upper()

    if status not in ["COMPLETE", "FURTHER ACTION"]:
        raise ValueError("Unexpected reply status: " + status)

    data["status"] = status
    return data


# saves replies that still need some action into the spreadsheet

def save_action(email, reply_data):
    excel_file = folder / "brokers_requiring_action.xlsx"

    if excel_file.exists():
        book = load_workbook(excel_file)
        sheet = book.active

        if sheet.cell(row=1, column=3).value is None:
            sheet.cell(row=1, column=3, value="Data Source")
    else:
        book = Workbook()
        sheet = book.active
        sheet.title = "Brokers"
        sheet.append(["Organisation", "Status", "Data Source"])

    organisation = (
        email.get("reply_to_email")
        or email.get("from_email")
        or email.get("from")
    )

    source = reply_data.get("datasource")
    if source is None:
        source = ""

    sheet.append([organisation, reply_data["status"], str(source)])
    book.save(excel_file)


# deals with one email at a time

def process_email(service, client, email):
    print("\nEmail from:", email["from"])
    print("Subject:", email["subject"])

    email_type = classify_email(client, email)
    print("Type:", email_type)

    if email_type == "Personal":
        print("No action needed")
        return

    if email_type == "Marketing":
        request = make_request(client, email)

        to = request.get("email_address")
        subject = request.get("subject")
        body = request.get("body")

        if not to or not subject or not body:
            raise ValueError("Missing information for the deletion request")

        print("Sending privacy request to:", to)
        send_email(service, to, subject, body)
        print("Sent")
        return

    if email_type == "Reply":
        reply_data = check_reply(client, email)

        print("Status:", reply_data["status"])
        print("Data source:", reply_data.get("datasource"))

        if reply_data["status"] == "FURTHER ACTION":
            save_action(email, reply_data)
            print("Saved to brokers_requiring_action.xlsx")
        else:
            print("Nothing else needed")


def main():
    service = gmail_login()

    # i used api_key in my original .env, but this also allows the normal name
    key = os.getenv("api_key") or os.getenv("OPENAI_API_KEY")
    if not key:
        raise RuntimeError("API key not found in .env")

    client = OpenAI(api_key=key)

    emails = get_unread(service)
    print("Unread emails:", len(emails))

    for item in emails:
        email_id = item["id"]

        try:
            email = get_email(service, email_id)
            process_email(service, client, email)

            # only mark it as read if everything above worked
            mark_read(service, email_id)

        except Exception as error:
            # failed emails stay unread so i can try them again
            print("Error processing email:", error)


if __name__ == "__main__":
    main()
